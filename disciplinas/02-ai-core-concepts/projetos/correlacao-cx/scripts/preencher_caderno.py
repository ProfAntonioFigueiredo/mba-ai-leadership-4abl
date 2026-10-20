"""Preenche o template 'Caderno de Hipóteses.docx' com a análise da base metrica_cx_corr_v2.csv.

Etapa 1 (hipóteses antes dos dados, SEM IA) e a identificação do grupo ficam em branco:
devem ser preenchidas pelo grupo.
"""
import copy
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

BASE = Path(__file__).parent.parent / "entrega"
doc = Document(BASE / "Caderno_de_Hipoteses_TEMPLATE.docx")
T = doc.tables


def set_cell(cell, text, size=9):
    """Substitui o conteúdo da célula mantendo o estilo do primeiro parágrafo."""
    paras = cell.paragraphs
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    p = paras[0]
    for r in p.runs:
        r._element.getparent().remove(r._element)
    for i, linha in enumerate(text.split("\n")):
        if i > 0:
            p = cell.add_paragraph(style=paras[0].style)
        p.add_run(linha).font.size = Pt(size)


def fill_rows(table, rows, size=8.5):
    """Preenche a partir da 2ª linha; clona a última linha se faltar espaço."""
    while len(table.rows) - 1 < len(rows):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    for r, valores in zip(table.rows[1:], rows):
        for c, v in zip(r.cells, valores):
            set_cell(c, v, size)


# 00 · Ferramentas de IA
set_cell(T[1].cell(0, 0),
         "Claude (Anthropic, via Claude Code no VS Code): interpretação da matriz e sugestão de hipóteses.\n"
         "Python (pandas, seaborn): cálculo da correlação de Pearson e Spearman e dos gráficos.")

# 02 · Olhar os dados
set_cell(T[3].cell(0, 0),
         "• 420 respostas; nenhum valor faltante.\n"
         "• Id: identificador, fica fora da matriz.\n"
         "• Status_Chamado = \"Resolvido\" em 100% das linhas: coluna que não varia, fica fora da matriz.\n"
         "• Age: 5 registros com idade 0 (Ids 17, 286, 297, 320, 413), provável erro de cadastro. Tratados como ausentes. "
         "Sem eles, a idade vai de 19 a 67 anos.\n"
         "• Tempo_Resolucao_h: de 1,1 h a 71,7 h. Média de 15,9 h, bem acima da mediana de 7,7 h (cauda longa de chamados demorados).\n"
         "• Distribuição: Baixa 52%, Média 30%, Alta 19% | Chat 40%, Telefone 34%, Portal 26%.\n"
         "• Canal e Criticidade são texto. Criticidade virou escala 1-2-3 e Canal virou 3 colunas 0/1.")

# 02 · Captura da matriz
c = T[4].cell(0, 0)
set_cell(c, "")
c.paragraphs[0].add_run().add_picture(str(BASE / "graficos" / "01_heatmap.png"), width=Cm(15))

# 02 · Correlações |r| > 0,4 (só leitura)
fill_rows(T[5], [
    ("e_NPS × RATINGS", "+0,89", "Quando a nota de satisfação sobe, a disposição em recomendar também sobe."),
    ("Criticidade × Tempo_Resolucao_h", "+0,89", "Quanto maior a criticidade, maior o tempo de resolução."),
    ("Age × Canal Telefone", "+0,74", "Quanto maior a idade, maior a presença no canal Telefone."),
    ("CES × Canal Telefone", "+0,65", "Atendimentos por telefone tendem a ter CES (esforço) mais alto."),
    ("Age × Canal Chat", "−0,62", "Quanto maior a idade, menor o uso do Chat."),
    ("Criticidade × RATINGS", "+0,59", "Quanto maior a criticidade, maior tende a ser a nota de satisfação."),
    ("CES × RATINGS", "−0,54", "Quando o esforço sobe, a satisfação tende a cair."),
    ("Criticidade × e_NPS", "+0,53", "Quanto maior a criticidade, maior tende a ser o e_NPS."),
    ("Age × CES", "+0,50", "Quanto maior a idade, maior tende a ser o esforço relatado."),
    ("e_NPS × CES", "−0,48", "Quando o esforço sobe, a disposição em recomendar tende a cair."),
    ("CES × Canal Chat", "−0,48", "Atendimentos por chat tendem a ter CES (esforço) mais baixo."),
    ("Tempo_Resolucao_h × RATINGS", "+0,42", "Quanto maior o tempo de resolução, maior tende a ser a nota (no agregado)."),
])

# 02 · Colunas fora da matriz
set_cell(T[6].cell(0, 0),
         "• Status_Chamado ficou fora porque é constante: só temos chamados resolvidos. Poderia explicar a satisfação de quem ficou sem solução "
         "(chamados abertos, cancelados ou reabertos). A pesquisa tem viés de seleção: só ouve quem foi atendido até o fim.\n"
         "• Id ficou fora por ser identificador. Se tivesse data, permitiria ver sazonalidade e tendência.\n"
         "• Canal e Criticidade só entram depois de convertidas em números. Os pares Canal × Canal (ex.: Chat × Telefone −0,58) "
         "foram ignorados, porque são negativos por construção (cada chamado tem um único canal).\n"
         "• Variáveis que NÃO existem e fazem falta: técnico ou nível (N1/N2/N3), SLA contratado, nº de transferências e reaberturas, "
         "área do colaborador. Elas poderiam explicar por que o crítico demora mais e é mais bem avaliado.\n"
         "• AUSÊNCIAS de correlação que também são informação: Criticidade × CES (0,01) e Tempo × CES (−0,02). "
         "O esforço do colaborador NÃO depende da gravidade nem da demora do chamado, e sim do canal. "
         "Criticidade × Canal (~0): chamados críticos não se concentram em nenhum canal.")

# 03 · Prompt
set_cell(T[7].cell(0, 0),
         "\"Você é um analista de dados. Aplique correlação estatística (Pearson) na base metrica_cx_corr_v2.csv "
         "(pesquisa interna de satisfação do Suporte de TI). Liste as correlações maiores que 0,4 ou menores que -0,4, "
         "interprete a matriz e sugira hipóteses para o diretor de Suporte.\"")

# 03 · Auditoria
fill_rows(T[8], [
    ("e_NPS e RATINGS têm correlação forte e positiva (0,89).", "Constatação",
     "Sim. Descreve o número. Com Spearman (mais adequado para escalas de 1 a 5) o valor se mantém: 0,90."),
    ("Chamados que demoram mais geram colaboradores mais satisfeitos (tempo × RATINGS = +0,42).", "Causal indevida",
     "Não. Separando por criticidade, a relação inverte: dentro de cada nível, mais tempo = MENOS satisfação (r de −0,49 a −0,57). "
     "É um Paradoxo de Simpson. A criticidade é a variável omitida."),
    ("Colaboradores mais velhos têm mais dificuldade: a idade aumenta o esforço (CES).", "Causal indevida",
     "Não. Controlando pelo canal, a idade quase não muda o CES (Telefone ≈ 3,7-3,9 em todas as faixas). "
     "A idade leva ao telefone, e o telefone é que tem mais esforço."),
    ("O canal Telefone está associado a maior esforço e menor satisfação.", "Hipótese",
     "Sim, como hipótese. CES médio 3,75 (vs. 2,2 no Chat) e RATINGS 2,4 (vs. 3,6). Falta medir espera, URA e transferências para explicar o porquê."),
    ("Chamados críticos são mais bem avaliados porque recebem técnicos mais seniores.", "Causal indevida",
     "Parcialmente. A correlação existe (+0,59), mas a base não tem a variável técnico ou nível. "
     "É uma hipótese plausível apresentada como fato. Precisa de dados de N1/N2/N3 para validar."),
    ("Chamados de alta criticidade demoram cerca de 10× mais (48 h vs. 4,5 h).", "Constatação",
     "Sim, é o número. Mas a IA não perguntou se isso cumpre o SLA. O esperado seria o crítico ser priorizado."),
])

# 03 · O que a IA deixou de questionar
set_cell(T[9].cell(0, 0),
         "• As 5 idades iguais a 0 (erro de cadastro), que distorcem as correlações com Age.\n"
         "• Que Status_Chamado é constante: a pesquisa só ouve chamados resolvidos (viés de seleção / sobrevivência).\n"
         "• Que as correlações Canal × Canal são artefato da codificação.\n"
         "• Que Pearson assume relação linear entre variáveis contínuas, e CES, RATINGS e e_NPS são escalas ordinais. "
         "Checamos com Spearman: tudo se mantém, exceto Tempo × RATINGS, que cai para 0,37 (abaixo do corte).\n"
         "• A variável omitida criticidade por trás de Tempo × Satisfação (Paradoxo de Simpson).\n"
         "• Se o tempo dos chamados críticos (48 h) está dentro do SLA.")

# 04 · Caderno final (coluna "Etapa 1" fica para o grupo)
fill_rows(T[10], [
    ("1", "Reduzir o tempo de resolução aumenta a satisfação em todos os níveis de criticidade.",
     "Tempo × RATINGS dentro de cada criticidade: −0,49 a −0,57 (no agregado, +0,42 enganoso)",
     "", "Criticidade (variável omitida que inverte o sinal: Paradoxo de Simpson)",
     "Piloto: fila expressa com meta de tempo menor para parte dos chamados de baixa criticidade vs. grupo controle; comparar RATINGS e e_NPS."),
    ("2", "O canal Telefone gera mais esforço e menos satisfação, independentemente da idade.",
     "CES × Telefone +0,65; CES × RATINGS −0,54; CES × Chat −0,48",
     "", "Idade (confundidor: 83% do público 46+ usa telefone). Controlado: o efeito persiste em todas as faixas",
     "Medir espera na URA, transferências e resolução no 1º contato; testar callback ou chat assistido para um grupo."),
    ("3", "A idade define o canal escolhido, e por isso a idade parece aumentar o esforço (correlação indireta).",
     "Age × Telefone +0,74; Age × Chat −0,62; Age × CES +0,50",
     "", "Espúria: Age × CES desaparece controlando o canal. Viés de seleção por canal",
     "Comparar o CES por faixa etária dentro do mesmo canal (feito: sem diferença relevante)."),
    ("4", "Chamados críticos recebem atendimento diferenciado (mais atenção e comunicação), o que eleva a satisfação.",
     "Criticidade × RATINGS +0,59; × e_NPS +0,53",
     "", "Nível do técnico e nº de interações (não estão na base); efeito alívio",
     "Cruzar com o nível do técnico (N1/N2/N3) e o nº de atualizações enviadas ao colaborador por chamado."),
    ("5", "Os chamados críticos demoram além do esperado (possível descumprimento de SLA).",
     "Criticidade × Tempo +0,89 (Alta = 48 h; Baixa = 4,5 h)",
     "", "Complexidade técnica e dependência de terceiros",
     "Comparar o tempo com o SLA por criticidade; mapear escalonamentos e fornecedores envolvidos."),
    ("6", "Satisfação (RATINGS) e recomendação (e_NPS) medem a mesma percepção.",
     "e_NPS × RATINGS +0,89 (Spearman 0,90)",
     "", "Mesma pessoa, mesmo momento (viés de resposta)",
     "Testar pesquisa enxuta em um piloto e ver se a informação se perde."),
])

# 04 · Derrubadas
set_cell(T[11].cell(0, 0),
         "• \"Demorar mais deixa o colaborador mais satisfeito\" (Tempo × RATINGS +0,42): derrubada. Dentro de cada criticidade a relação é negativa. "
         "Com Spearman, ainda cai para 0,37, abaixo do corte.\n"
         "• \"Colaboradores mais velhos têm mais dificuldade\" (Age × CES +0,50): derrubada. No mesmo canal, o esforço é praticamente igual "
         "em todas as faixas etárias. O efeito é do canal Telefone, não da idade.")

# 05 · Recomendação
set_cell(T[13].cell(0, 0),
         "Tratar o tempo de resolução como alavanca de satisfação: fixar metas de redução por nível de criticidade e começar pelo canal Telefone, "
         "que concentra o maior esforço e o público 46+.")
set_cell(T[14].cell(0, 0),
         "Dentro de cada criticidade, mais tempo = menos satisfação (r de −0,49 a −0,57); o Telefone tem CES de 3,75 contra 2,24 do Chat, em qualquer idade. "
         "Não prova causa: é uma pesquisa observacional, só com chamados resolvidos, sem dados de técnico, SLA ou transferências.")
set_cell(T[15].cell(0, 0),
         "Piloto de 8 semanas: sortear metade dos chamados de baixa e média criticidade para uma fila expressa (meta de tempo 50% menor) e, no Telefone, "
         "oferecer callback. Comparar RATINGS, e_NPS e CES com o grupo controle. Diferença significativa indica causa.")
set_cell(T[16].cell(0, 0),
         "Se agirmos e for falsa: custo de reorganizar a fila e o esforço da equipe, sem ganho de satisfação (reversível, piloto limitado). "
         "Se não agirmos e for verdadeira: colaboradores parados por mais tempo (perda de produtividade), queda contínua de e_NPS e "
         "a experiência pior concentrada no público 46+ que depende do telefone.")

out = BASE / "Caderno_de_Hipoteses_CX_preenchido.docx"
doc.save(out)
print("salvo:", out)
