"""Convert course documents into a traceable Markdown library (never changes sources).

python3 tools/ingest_library.py --source /path/to/MBAAILEADERSHIP
Requires pandoc, Poppler, python-pptx and openpyxl. Cloud placeholders are reported,
not opened: materialize them before rerunning. Identical inputs share one document.
"""
from pathlib import Path
import argparse, hashlib, json, os, re, subprocess, unicodedata

ROOT = Path(__file__).resolve().parents[1]
EXTS = {'.docx', '.pdf', '.pptx', '.xlsx', '.md', '.txt', '.html', '.ipynb', '.png', '.jpg', '.jpeg', '.webp'}
SKIP = {'.git', '.venv', 'node_modules', 'dist', '__pycache__', '.vscode',
        'test-results', 'playwright-report', 'mba-ai-leadership-4abl', ROOT.name}

def slug(value):
    value = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', value).strip('-')[:85] or 'documento'

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def available(path):
    return path.is_file() and not (getattr(path.stat(), 'st_flags', 0) & 0x40000000)

def run(*args, cwd=None):
    return subprocess.check_output(args, cwd=cwd, text=True, stderr=subprocess.PIPE, timeout=120)

def discipline(rel):
    if rel.startswith(('Aulas/AI-Core', 'AI Core Concepts/')): return '02'
    if rel.startswith('AI Sociedades'): return '03'
    if rel.startswith('Vibe Coding'): return '21'
    return 'geral'

def table(rows):
    def cell(v): return str(v if v is not None else '').replace('|', '\\|').replace('\n', '<br>')
    rows = [[cell(c) for c in row] for row in rows]
    if not rows: return ''
    width = max(map(len, rows)); rows = [row + [''] * (width-len(row)) for row in rows]
    return '\n'.join(['| '+' | '.join(rows[0])+' |', '| '+' | '.join(['---']*width)+' |',
                      *['| '+' | '.join(row)+' |' for row in rows[1:]]])+'\n'

def extract(src, folder):
    ext = src.suffix.lower(); note = ''
    if ext == '.docx':
        body = run('pandoc', str(src), '-t', 'gfm', '--wrap=none', '--extract-media=assets', cwd=folder)
        note = 'Conversão do DOCX com títulos, listas, tabelas e imagens incorporadas. A paginação original não é preservada.'
    elif ext == '.pdf':
        raw = run('pdftotext', '-layout', str(src), '-')
        pages = raw.split('\f')
        (folder/'assets').mkdir(exist_ok=True)
        run('pdftoppm','-jpeg','-scale-to','1400','-jpegopt','quality=78',str(src),'assets/pagina',cwd=folder)
        parts=[]
        for i,image in enumerate(sorted((folder/'assets').glob('pagina-*.jpg')),1):
            text=pages[i-1].rstrip() if i<=len(pages) else ''
            parts.append('## Página '+str(i)+'\n\n![Página '+str(i)+' do documento]('+image.relative_to(folder).as_posix()+')\n\n'+('```text\n'+text+'\n```' if text else '*Página sem texto extraível; consulte a imagem acima.*'))
        body='\n\n'.join(parts)
        note = 'Texto extraído e imagem de cada página para preservar a consulta a diagramas e figuras. Não é resumo nem OCR; texto presente apenas nas imagens não participa da busca.'
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(src); parts = []
        for i, slide in enumerate(prs.slides, 1):
            parts.append('## Slide '+str(i))
            for shape in slide.shapes:
                if shape.has_text_frame: parts.append(shape.text)
                if shape.has_table: parts.append(table([[c.text for c in row.cells] for row in shape.table.rows]))
                if shape.shape_type == 13:
                    asset=folder/'assets'/f'slide-{i}-{shape.shape_id}.{shape.image.ext}'
                    asset.parent.mkdir(exist_ok=True);asset.write_bytes(shape.image.blob)
                    parts.append(f'![Imagem do slide {i}]({asset.relative_to(folder).as_posix()})')
            if slide.has_notes_slide:
                notes=slide.notes_slide.notes_text_frame.text
                if notes.strip():parts.append('### Notas do apresentador\n\n'+notes)
        body='\n\n'.join(parts);note='Texto, tabelas, imagens incorporadas e notas por slide. Layout, gráficos e SmartArt podem exigir consulta ao original.'
    elif ext == '.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(src, data_only=False, read_only=True)
        body='\n\n'.join('## '+ws.title+'\n\n'+table(list(ws.values)) for ws in wb.worksheets)
        wb.close();note='Células e fórmulas exportadas por planilha. Não executa fórmulas; gráficos e formatação não são reproduzidos.'
    elif ext in {'.png','.jpg','.jpeg','.webp'}:
        asset=folder/'assets'/('original'+ext);asset.parent.mkdir(exist_ok=True)
        asset.write_bytes(src.read_bytes())
        body='!['+src.stem+']('+asset.relative_to(folder).as_posix()+')'
        note='Imagem original incorporada para consulta visual; o texto da imagem não foi transcrito por OCR.'
    elif ext == '.ipynb':
        nb=json.loads(src.read_text());parts=[]
        for cell in nb.get('cells',[]):
            content=''.join(cell.get('source',[]))
            parts.append(content if cell.get('cell_type')=='markdown' else '```python\n'+content+'\n```')
            for output in cell.get('outputs',[]):
                text=''.join(output.get('text',[]) or output.get('data',{}).get('text/plain',[]))
                if text:parts.append('```text\n'+text+'\n```')
        body='\n\n'.join(parts);note='Células e saídas textuais já gravadas. O notebook não foi executado na conversão.'
    elif ext == '.html':
        body=run('pandoc', str(src), '-f', 'html', '-t', 'gfm', '--wrap=none')
        note='Conteúdo textual do HTML; código interativo e scripts não são executados.'
    else:
        body=src.read_text(encoding='utf-8-sig')
        note='Markdown preservado.' if ext=='.md' else 'Transcrição textual integral; não é resumo.'
        if ext=='.txt':body='```text\n'+body+'\n```'
    if not body.strip():raise ValueError('Conversão sem conteúdo')
    return body, note

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True, type=Path)
    parser.add_argument('--fallback', action='append', default=[], type=Path,
                        help='Additional source root with the same relative paths')
    args=parser.parse_args();source=args.source.resolve()
    if not source.is_dir():parser.error('A pasta de origem não existe: '+str(source))
    if source==ROOT or ROOT in source.parents:parser.error('Use a pasta do acervo original, fora da biblioteca gerada.')
    imported=json.loads((ROOT/'catalogo/importacao.json').read_text())
    lookup={e['origem']:e for e in imported}
    existing=json.loads((ROOT/'catalogo/biblioteca.json').read_text()) if (ROOT/'catalogo/biblioteca.json').exists() else []
    prior={origin:e for e in existing for origin in e['origens']}
    lib=ROOT/'biblioteca';lib.mkdir(exist_ok=True)
    entries=[];groups={};pending=[]
    for base,dirs,files in os.walk(source):
        dirs[:]=sorted(d for d in dirs if d not in SKIP and not d.startswith('.'))
        for name in sorted(files):
            original=Path(base)/name
            if original.suffix.lower() not in EXTS:continue
            rel=original.relative_to(source).as_posix()
            # Code repositories remain intact; their Markdown is indexed separately by build_library.
            if rel.startswith(('Vibe Coding Executivo/speakup-lovable/', 'Vibe Coding Executivo/speakup-sdd/')) and original.suffix.lower()!='.txt':continue
            candidates=[original]
            if rel in lookup:candidates.append(ROOT/lookup[rel]['destino'])
            candidates.append(source/'mba-ai-leadership-4abl/acervo-local/originais'/rel)
            candidates += [p/rel for p in args.fallback]
            if rel.startswith('Aulas/'):
                candidates += [p/rel.removeprefix('Aulas/') for p in args.fallback]
            src=next((p for p in candidates if available(p)),None)
            old=prior.get(rel)
            if src is None:
                if old and (ROOT/old['markdown']).is_file():
                    if old['sha256'] not in groups:entries.append(old.copy());groups[old['sha256']]=entries[-1]
                    continue
                pending.append({'origem':rel,'motivo':'Fonte somente na nuvem ou indisponível'});continue
            try:
                h=digest(src)
                # Adapted Markdown may be read from the verified organized copy.
                if rel in lookup and src==(ROOT/lookup[rel]['destino']) and h==lookup[rel]['sha256_destino']:
                    h=lookup[rel]['sha256']
                if h in groups:
                    if rel not in groups[h]['origens']:groups[h]['origens'].append(rel)
                    continue
                if original.suffix.lower()=='.md' and rel in lookup:
                    target=ROOT/lookup[rel]['destino']
                    entry={'titulo':original.stem,'disciplina':discipline(rel),'origens':[rel],
                           'sha256':h,'formato':'.md','markdown':target.relative_to(ROOT).as_posix(),
                           'observacao':'Versão Markdown existente reutilizada, sem cópia adicional.',
                           'sha256_markdown':digest(target)}
                    entries.append(entry);groups[h]=entry;continue
                folder=lib/discipline(rel)/(slug(original.stem)+'-'+h[:8]);folder.mkdir(parents=True,exist_ok=True)
                body,note=extract(src.resolve(),folder)
                # GitHub-safe Markdown; relative media stays beside each document.
                title=original.stem
                target=folder/'conteudo.md'
                attribution='Material do acervo acadêmico fornecido pelo usuário. Autoria e direitos permanecem com os titulares indicados na fonte.'
                target.write_text('# '+title+'\n\n[Índice da biblioteca](../../README.md)\n\n'
                    +'> Fonte: `'+rel+'`  \n> SHA-256: `'+h+'`  \n> '+note+'  \n> '+attribution+'\n\n---\n\n'+body+'\n')
                entry={'titulo':title,'disciplina':discipline(rel),'origens':[rel],'sha256':h,
                       'formato':original.suffix.lower(),'markdown':target.relative_to(ROOT).as_posix(),
                       'observacao':note,'sha256_markdown':digest(target)}
                entries.append(entry);groups[h]=entry
            except (OSError,ValueError,subprocess.SubprocessError,UnicodeError) as exc:
                pending.append({'origem':rel,'motivo':str(exc)[:250]})
    entries.sort(key=lambda e:(e['disciplina'],e['titulo'].casefold()))
    (ROOT/'catalogo/biblioteca.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2)+'\n')
    (ROOT/'catalogo/pendencias-conversao.json').write_text(json.dumps(pending,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'documentos':len(entries),'origens_convertidas':sum(len(e['origens']) for e in entries),'pendentes':len(pending)},ensure_ascii=False))

if __name__=='__main__':main()
